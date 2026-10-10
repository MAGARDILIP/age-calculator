from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] # Title slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Age Calculator Web Application"
    subtitle.text = "Formative Assessment 2 (FA-2) Final Presentation\nDeveloped with React.js"

    # 2. Project Overview
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Project Overview"
    tf = content.text_frame
    tf.text = "Objective: Build an interactive, aesthetic Age Calculator as a Single Page Application (SPA)."
    p = tf.add_paragraph()
    p.text = "Accurate exact age calculation (Years, Months, Days)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Modern Glassmorphism UI with Dark/Light Theme Support."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Integration with external REST APIs for historical fun facts."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Persistent user profiles without requiring a database."
    p.level = 1

    # 3. Technologies Used
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Technologies Used"
    tf = content.text_frame
    tf.text = "Frontend Framework: React.js (Vite)"
    p = tf.add_paragraph()
    p.text = "Styling: Pure CSS with CSS Variables for Theming"
    p = tf.add_paragraph()
    p.text = "State Management: React Hooks (useState, useEffect, Context API)"
    p = tf.add_paragraph()
    p.text = "Data Persistence: Web LocalStorage API"
    p = tf.add_paragraph()
    p.text = "Deployment & Hosting: Vercel"

    # 4. Key Feature: UI & Theme System
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "UI/UX & Theming System"
    tf = content.text_frame
    tf.text = "Built a custom Dark/Light toggle using Context API."
    p = tf.add_paragraph()
    p.text = "Global CSS Variables control the color palette seamlessly."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Implemented 'Glassmorphism' effects (blurred backgrounds, translucent cards)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Fully responsive layout adapting to mobile and desktop displays."
    p.level = 1

    # 5. Key Feature: Wikipedia API Integration
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "API Integration: Historical Facts"
    tf = content.text_frame
    tf.text = "Dynamic content fetching based on the user's birth year."
    p = tf.add_paragraph()
    p.text = "Replaced flaky HTTP APIs with the Wikipedia REST API for reliable, HTTPS-native data."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Graceful fallbacks: If fetching fails, the app generates mathematical facts algorithmically (e.g., Leap Year checks)."
    p.level = 1

    # 6. Key Feature: Profile Management
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Profile Management (CRUD)"
    tf = content.text_frame
    tf.text = "Users can Create, Read, Update, and Delete saved age profiles."
    p = tf.add_paragraph()
    p.text = "Migrated from json-server to LocalStorage for complete serverless architecture."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Allows the application to run as a static SPA on Vercel without a backend."
    p.level = 1

    # 7. Challenges & Solutions
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Challenges & Solutions"
    tf = content.text_frame
    tf.text = "Mixed Content Errors on Vercel"
    p = tf.add_paragraph()
    p.text = "Solution: Migrated API usage to an HTTPS-supported Wikipedia endpoint."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "CSS Variable Inheritance Bugs in Light Theme"
    p = tf.add_paragraph()
    p.text = "Solution: Correctly scoped root color variables to the app-wrapper."
    p.level = 1

    # 8. Conclusion
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    content = slide.placeholders[1]
    title.text = "Conclusion & Demonstration"
    tf = content.text_frame
    tf.text = "Successfully delivered a robust, fast, and aesthetically pleasing React application."
    p = tf.add_paragraph()
    p.text = "Meets all FA-2 requirements including API usage and state management."
    p = tf.add_paragraph()
    p.text = "Live Demo: Available on Vercel."
    p = tf.add_paragraph()
    p.text = "Thank You! Any Questions?"

    # Save presentation
    output_path = "Age_Calculator_FA2_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation generated successfully at {output_path}")

if __name__ == '__main__':
    create_presentation()
