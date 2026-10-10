import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def add_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = RGBColor(41, 128, 185) # A nice professional blue

def main():
    doc = Document()

    # Title Page
    title = doc.add_heading('Age Calculator — Comprehensive Project Report (FA-2)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('\n' * 5)
    info = doc.add_paragraph('Course: Advanced Web Technologies (AWT)\nDepartment: MCA, Semester III (SYMCA)\nAcademic Year: 2026–27\nCollege: Pimpri Chinchwad College of Engineering (PCCOE)\n\nSubmitted by:\nMagar Dilip Balu (PRN: 125M1H055)\nRoll No: 52')
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # Project Details Table
    add_heading(doc, 'Project Information', 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text, hdr_cells[1].text = 'Field', 'Details'
    
    details = [
        ('Project Title', 'Age Calculator & Profile Manager'),
        ('Student Name', 'Magar Dilip Balu'),
        ('PRN', '125M1H055'),
        ('Course Owner', 'Dr. Avinash Chormale'),
        ('Module Coordinator', 'Dr. Rajkamal Sangole'),
        ('Assessment', 'Formative Assessment-II (FA-2)'),
        ('Deployment Environment', 'Vercel (Serverless Edge Network)'),
        ('Live URL', 'https://age-calculator-ivory-gamma.vercel.app/'),
        ('GitHub Repository', 'https://github.com/MAGARDILIP/age-calculator.git')
    ]
    for k, v in details:
        row_cells = table.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = v

    doc.add_page_break()

    # Section 1: Introduction & Objectives
    add_heading(doc, '1. Introduction & Objectives', 1)
    doc.add_paragraph('1.1 Background\nAge calculation is a universal requirement, from filling out application forms to determining eligibility criteria. The Age Calculator web application was initially developed in Formative Assessment-I (FA-1) to provide users with a fast, accurate, and aesthetically pleasing way to calculate their exact age in years, months, and days. It also featured additional fun metrics like a countdown to the next birthday and total days lived.')
    
    doc.add_paragraph('1.2 FA-2 Enhancements & Scope\nWhile Phase-I established the core logic and visual design, Phase-II (FA-2) aimed to evolve the project from a simple utility into a full-fledged, robust Single Page Application (SPA). The primary objectives of FA-2 included:')
    objectives = [
        'Implementing Client-Side Routing: Allowing navigation between multiple distinct pages (Calculator, Dashboard, Add/Edit Profile) without page reloads.',
        'Data Persistence (Serverless CRUD): Implementing a system to save, view, edit, and delete user profiles directly within the browser using the Web LocalStorage API, removing the reliance on a local JSON server and allowing seamless cloud deployment.',
        'External API Integration: Fetching dynamic, historical fun facts based on the user\'s birth year from a reliable public API.',
        'Advanced React Concepts: Deepening the usage of React Hooks (useEffect, useCallback, useState, useContext) to manage complex component lifecycles and state.',
        'Testing & QA: Establishing a testing suite using Vitest to ensure mathematical accuracy of the core algorithms.',
        'Production Deployment: Successfully deploying the application to the internet via Vercel.'
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    doc.add_page_break()

    # Section 2: System Architecture & Technology Stack
    add_heading(doc, '2. System Architecture & Technology Stack', 1)
    doc.add_paragraph('2.1 High-Level Architecture\nThe Age Calculator is designed as a modular React Single Page Application (SPA). It uses React Router for seamless view transitions. The application state is decoupled into local component state (for forms and immediate UI feedback) and global context (for the application theme). Data fetching is isolated into dedicated service layers (`api.js` for external APIs, `profileService.js` for local data persistence).')
    
    doc.add_paragraph('2.2 Technology Stack')
    tech_table = doc.add_table(rows=1, cols=3)
    tech_table.style = 'Table Grid'
    t_hdr = tech_table.rows[0].cells
    t_hdr[0].text, t_hdr[1].text, t_hdr[2].text = 'Category', 'Technology', 'Purpose'
    
    techs = [
        ('Frontend Library', 'React.js (v19)', 'Declarative UI creation and component lifecycle management.'),
        ('Build Tool', 'Vite', 'Ultra-fast Hot Module Replacement (HMR) and optimized production bundling.'),
        ('Routing', 'React Router DOM v7', 'Client-side routing across different application views.'),
        ('Styling', 'Pure CSS3 & CSS Variables', 'Implementing the glassmorphism design system and dynamic theme toggling without heavy CSS frameworks.'),
        ('Testing', 'Vitest & JSDOM', 'Executing isolated unit tests and component rendering tests.'),
        ('Data Persistence', 'Web LocalStorage API', 'Storing user profiles in a serverless environment (CRUD operations).'),
        ('External API', 'Wikipedia REST API', 'Fetching reliable, HTTPS-native historical events and summaries.'),
        ('Deployment', 'Vercel', 'Hosting the compiled SPA on a global edge network.')
    ]
    for cat, tech, purp in techs:
        row = tech_table.add_row().cells
        row[0].text, row[1].text, row[2].text = cat, tech, purp

    doc.add_page_break()

    # Section 3: Detailed Implementation of React Concepts
    add_heading(doc, '3. Detailed Implementation of React Concepts', 1)
    
    doc.add_paragraph('3.1 Functional Components & Hooks')
    doc.add_paragraph('The entire application is built using React functional components, strictly adhering to modern React paradigms. Various hooks were utilized to manage state and side effects:')
    hooks = [
        'useState: Used extensively to manage form inputs, toggle states, loading indicators, and error messages across all components.',
        'useEffect: Crucial for executing side effects. For example, it is used in the `FunFacts` component to trigger an asynchronous fetch to the Wikipedia API whenever the user\'s birth year changes. It is also used to initialize the theme from LocalStorage on first load.',
        'useCallback: Implemented in the `Home.jsx` component to memoize the `handleCalculate` function. This prevents the function from being recreated on every render, optimizing performance when passed as a prop to child components.',
        'useContext: A custom `ThemeContext` was created to provide the current theme (light or dark) and a toggle function globally to the entire application tree without prop drilling.'
    ]
    for h in hooks:
        doc.add_paragraph(h, style='List Bullet')

    doc.add_paragraph('\n3.2 Client-Side Routing')
    doc.add_paragraph('Using React Router DOM, the application avoids full page reloads. The routing architecture is defined in `App.jsx` with the following routes:')
    routes = [
        '/ : The main calculator interface.',
        '/profiles : The dashboard listing all saved profiles.',
        '/profiles/new : A form to create a new profile.',
        '/profiles/edit/:id : The same form pre-populated to edit an existing profile, utilizing URL parameters (`useParams`).'
    ]
    for r in routes:
        doc.add_paragraph(r, style='List Bullet')

    doc.add_page_break()

    # Section 4: Data Persistence & CRUD Operations
    add_heading(doc, '4. Data Persistence (Serverless CRUD)', 1)
    doc.add_paragraph('4.1 Migration to LocalStorage\nInitially, the project utilized a local JSON server (`json-server`) to mock a REST API. However, upon deploying to Vercel, it became evident that Vercel\'s serverless, immutable environment does not support persistent local file modifications. To ensure the application functioned perfectly in the cloud, the entire data layer was rewritten to utilize the browser\'s native LocalStorage API.')
    
    doc.add_paragraph('4.2 The profileService.js Architecture\nA dedicated service file (`src/services/profileService.js`) abstracts all LocalStorage interactions, providing a clean, Promise-based API that mimics a real backend. This includes:')
    crud = [
        'Create (addProfile): Generates a unique ID using `Date.now()`, serializes the profile object to a JSON string, and appends it to the stored array.',
        'Read (getProfiles / getProfileById): Deserializes the JSON string from LocalStorage and returns the array of profiles or a specific profile by ID.',
        'Update (updateProfile): Maps over the existing array, replacing the modified profile, and reserializes the array.',
        'Delete (deleteProfile): Filters out the targeted ID and saves the mutated array back to LocalStorage.'
    ]
    for c in crud:
        doc.add_paragraph(c, style='List Bullet')

    doc.add_page_break()

    # Section 5: API Integration & Data Fetching
    add_heading(doc, '5. API Integration & Data Fetching', 1)
    doc.add_paragraph('5.1 Wikipedia API Integration\nTo fulfill the requirement of displaying historical fun facts based on the user\'s birth year, the application integrates with the Wikipedia REST API (`https://en.wikipedia.org/api/rest_v1/page/summary/`).')
    
    doc.add_paragraph('5.2 Overcoming CORS and Mixed-Content Challenges\nInitially, the Numbers API (`http://numbersapi.com`) was used. However, this API only supported HTTP, which resulted in fatal "Mixed Content" security blocks when the application was deployed to Vercel (which enforces HTTPS). Attempts to use free CORS proxies (like corsproxy.io or allorigins.win) resulted in severe rate-limiting and Cloudflare 522 timeouts.')
    doc.add_paragraph('The solution was to completely migrate to the Wikipedia API, which natively supports HTTPS and possesses extremely high reliability, requiring no flaky proxies.')

    doc.add_paragraph('5.3 Algorithmic Graceful Fallback\nIn the event that the Wikipedia API fails or the user loses internet connection, the `fetchYearFact` function features a graceful fallback mechanism. It catches the network error and algorithmically computes whether the user\'s birth year was a leap year or a common year, returning a friendly, mathematically accurate string instead of a crash or blank screen.')

    doc.add_page_break()

    # Section 6: UI/UX & Theming System
    add_heading(doc, '6. UI/UX & Theming System', 1)
    doc.add_paragraph('6.1 Glassmorphism Aesthetic\nThe application employs a "Glassmorphism" design language, characterized by frosted-glass effects, blurred backgrounds (`backdrop-filter: blur()`), and subtle translucent borders. This modern aesthetic is achieved entirely through pure CSS3, keeping the bundle size small.')

    doc.add_paragraph('6.2 CSS Variable Inheritance (Dark & Light Themes)\nTheming is strictly managed via CSS variables defined on the `:root` and `[data-theme="light"]` selectors. The `App.jsx` component applies a `data-theme` attribute to the main wrapper div, which cascaded the CSS variables to all child components.')
    
    doc.add_paragraph('A notable challenge was CSS inheritance bugs where the `body` tag retained the dark theme text color in light mode, causing text to become invisible (white-on-white). This was resolved by explicitly scoping the `color: var(--text-primary)` declaration to the `.app-wrapper` class, ensuring perfect transitions between Dark and Light modes.')

    doc.add_page_break()

    # Section 7: Testing Strategy
    add_heading(doc, '7. Testing Strategy', 1)
    doc.add_paragraph('Quality assurance is enforced through an automated testing suite utilizing Vitest and React Testing Library.')

    tests = [
        'Utility Testing: The core `calculateAge` function is tested with various edge cases, including leap years, future dates (which should return null), and exact birthday matches.',
        'Component Testing: The `AgeResult.jsx` component is mounted in a simulated JSDOM environment to ensure it renders the calculated metrics (Years, Months, Days) correctly into the DOM.',
        'Continuous Integration Readiness: The tests run instantly via `npm run test`, ensuring no regressions are introduced during future development.'
    ]
    for t in tests:
        doc.add_paragraph(t, style='List Bullet')

    doc.add_page_break()

    # Section 8: Deployment Architecture
    add_heading(doc, '8. Deployment Architecture (Vercel)', 1)
    doc.add_paragraph('The application is continuously deployed to Vercel. The deployment process involved several critical configurations:')
    
    doc.add_paragraph('8.1 vercel.json Configuration')
    doc.add_paragraph('Because React Router uses the HTML5 History API for client-side routing, directly navigating to a route like `/profiles` on a static host typically results in a 404 error. A `vercel.json` file was created with the following configuration to rewrite all requests to `index.html`:')
    
    code = doc.add_paragraph('{\n  "rewrites": [\n    { "source": "/(.*)", "destination": "/index.html" }\n  ]\n}')
    code.style = 'Intense Quote'

    doc.add_paragraph('8.2 Build Pipeline')
    doc.add_paragraph('Vercel automatically detects the Vite framework, runs `npm install`, executes `npm run build`, and serves the optimized static assets from the `dist` directory across its global Edge Network, ensuring blazing fast load times worldwide.')

    doc.add_page_break()

    # Section 9: Advanced Use Cases & User Journeys
    add_heading(doc, '9. Advanced Use Cases & User Journeys', 1)
    journeys = [
        'UC-1 (Calculate Age): User lands on the Home page, selects a birth date, and instantly sees their age breakdown and next birthday countdown.',
        'UC-2 (Theme Toggle): User clicks the theme icon in the header. The app smoothly transitions to Light Mode, and this preference is saved.',
        'UC-3 (Save Profile): User navigates to the Add Profile page, enters a name, date, and zodiac sign. The form validates the input and saves it to LocalStorage.',
        'UC-4 (Search & Filter): User visits the Dashboard, which is populated with profiles. The user types "John" in the search bar, and the grid instantly filters the results in real-time.',
        'UC-5 (Edit/Delete): User clicks "Edit" on a profile card, is routed to the form pre-populated with data, updates it, and saves. The user then deletes a different profile, prompting a re-render of the remaining grid.'
    ]
    for j in journeys:
        doc.add_paragraph(j, style='List Bullet')

    doc.add_page_break()

    doc.add_heading('10. Conclusion & Future Scope', 1)
    doc.add_paragraph('The Age Calculator (FA-2) project successfully transforms a simple mathematical utility into a highly interactive, beautifully designed, and structurally sound React application. By leveraging modern React Hooks, CSS Variables, client-side routing, and serverless data persistence, the project strictly adheres to modern software engineering best practices.')
    
    p = doc.add_paragraph()
    run = p.add_run('Future Scope:')
    run.bold = True
    
    future = [
        'Cloud Sync: Implementing Firebase or Supabase to allow users to sync profiles across multiple devices.',
        'Push Notifications: Integrating Web Push API to send browser notifications when a saved profile\'s birthday is approaching.',
        'Export Functionality: Adding the ability to download the calculated age report as a PDF or share it directly to social media.'
    ]
    for f in future:
        doc.add_paragraph(f, style='List Bullet')

    doc.add_paragraph('\n\n--- End of Report ---')
    doc.add_paragraph('Submitted by: Magar Dilip Balu (PRN: 125M1H055)')

    output_file = 'Progress_Report_FA2.docx'
    doc.save(output_file)
    print(f"Successfully generated extended {output_file}")

if __name__ == '__main__':
    main()
