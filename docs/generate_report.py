import os
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def main():
    doc = Document()

    # Title
    title = doc.add_heading('Age Calculator — Progress Report (FA-1)', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph('Course: Advanced Web Technologies (AWT)\nDepartment: MCA, Semester III (SYMCA)\nAcademic Year: 2026–27\nCollege: Pimpri Chinchwad College of Engineering (PCCOE)')

    # Info Table
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Field'
    hdr_cells[1].text = 'Details'
    
    details = [
        ('Student Name', 'Magar Dilip Balu'),
        ('PRN', '125M1H055'),
        ('Roll No.', '52'),
        ('Project Title', 'Age Calculator'),
        ('Course Owner', 'Dr. Avinash Chormale'),
        ('Module Coordinator', 'Dr. Rajkamal Sangole'),
        ('Assessment', 'Formative Assessment-I (FA-1)')
    ]
    for k, v in details:
        row_cells = table.add_row().cells
        row_cells[0].text = k
        row_cells[1].text = v


    # Section 1
    doc.add_heading('1. Problem Statement', level=1)
    doc.add_paragraph('People often need to know their exact age in years, months, and days for various purposes such as official document applications, visa processing, competitive exam eligibility, insurance calculations, and retirement planning. Manually computing age across leap years, varying month lengths, and calendar boundaries is error-prone and tedious.')
    doc.add_paragraph('Proposed Solution: A modern, interactive Age Calculator web application built using React.js that allows users to input their date of birth and instantly view their exact age along with detailed breakdowns, birthday countdown, and fun birth-date facts.', style='List Bullet')

    # Section 2
    doc.add_heading('2. Objectives', level=1)
    objs = [
        'Calculate and display exact age in years, months, and days from a given date of birth.',
        'Provide a detailed age breakdown in total months, weeks, days, hours, and minutes.',
        'Show a countdown to the user\'s next birthday.',
        'Display fun facts such as zodiac sign, birth day name, generation, and birth season.',
        'Maintain a session-based history of all calculations performed.',
        'Support dark and light themes with a responsive, mobile-friendly user interface.'
    ]
    for obj in objs:
        doc.add_paragraph(obj, style='List Number')

    # Section 3
    doc.add_heading('3. Use Case Identification', level=1)
    ucs = [
        'UC-1: Input Date of Birth: The user selects their date of birth using a date picker.',
        'UC-2: Calculate Age: The system calculates the exact age in years, months, and days based on the provided date of birth.',
        'UC-3: View Detailed Breakdown: The user views their age broken down into total months, weeks, days, hours, and minutes.',
        'UC-4: View Next Birthday: The user sees a countdown showing the exact number of days remaining until their next birthday.',
        'UC-5: View Fun Facts: The user discovers interesting facts about their birth date (e.g., Zodiac sign, birth season).',
        'UC-6: Manage History: The user views a history of their past age calculations and can clear the history.',
        'UC-7: Toggle Theme: The user switches between light and dark modes for a personalized viewing experience.'
    ]
    for uc in ucs:
        doc.add_paragraph(uc, style='List Bullet')

    # Section 4
    doc.add_heading('4. UI Design & Wireframe', level=1)
    doc.add_paragraph('The application follows a modern glassmorphism design aesthetic with:')
    designs = [
        'Dark mode as default theme (with light mode toggle)',
        'Purple-to-cyan gradient accents',
        'Inter font (Google Fonts) for clean typography',
        'Frosted glass card effects with subtle hover animations',
        'Fully responsive grid layout (2-column on desktop → single column on mobile)'
    ]
    for d in designs:
        doc.add_paragraph(d, style='List Bullet')
    
    if os.path.exists('screenshots/wireframe.jpg'):
        doc.add_picture('screenshots/wireframe.jpg', width=Inches(2.5))
        p = doc.add_paragraph('Figure 1: UI Wireframe Layout')
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


    # Section 5
    doc.add_heading('5. Functional Requirements', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'ID', 'Requirement', 'Status'
    frs = [
        ('FR-1', 'Date of birth input via native date picker', 'Implemented'),
        ('FR-2', 'Exact age calculation (Years / Months / Days)', 'Implemented'),
        ('FR-3', 'Detailed breakdown (total months, weeks, days, hours, minutes)', 'Implemented'),
        ('FR-4', 'Next birthday countdown with days remaining', 'Implemented'),
        ('FR-5', 'Calculation history with clear functionality', 'Implemented'),
        ('FR-6', 'Dark/Light theme toggle', 'Implemented'),
        ('FR-7', 'Responsive layout (mobile, tablet, desktop)', 'Implemented'),
        ('FR-8', 'Input validation with error messages', 'Implemented'),
        ('FR-9', 'Fun facts (zodiac sign, birth day, season, generation, leap year)', 'Implemented')
    ]
    for id_str, req, status in frs:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = id_str, req, status

    # Section 6
    doc.add_heading('6. Technology Stack', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = 'Technology', 'Purpose'
    techs = [
        ('React.js', 'Component-based UI development'),
        ('JavaScript (ES6+)', 'Application logic, date calculations'),
        ('Vite', 'Build tool and development server'),
        ('HTML5', 'Semantic page structure'),
        ('CSS3', 'Styling with custom properties, glassmorphism, responsive design'),
        ('VS Code', 'Development environment')
    ]
    for t, p in techs:
        row = table.add_row().cells
        row[0].text, row[1].text = t, p

    # Section 7
    doc.add_heading('7. React Concepts Implemented', level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text = 'Concept', 'Components Using It'
    react_concepts = [
        ('JSX', 'All 10 components'),
        ('Functional Components', 'Header, DateInput, AgeResult, AgeBreakdown, BirthdayCountdown, FunFacts, HistoryItem, HistoryList, Footer'),
        ('Props', 'Header, DateInput, AgeResult, AgeBreakdown, BirthdayCountdown, FunFacts, HistoryItem, HistoryList'),
        ('useState', 'App (theme, result, breakdown, birthday, funFacts, history), DateInput (inputDate, error)'),
        ('Event Handling', 'DateInput (onChange, onSubmit), Header (theme toggle), HistoryList (clear)'),
        ('Conditional Rendering', 'Results section, Error messages, Birthday today, Empty history state, Theme icon'),
        ('Lists & Keys', 'AgeBreakdown (stat items), FunFacts (facts list), HistoryList (history entries)')
    ]
    for c, comp in react_concepts:
        row = table.add_row().cells
        row[0].text, row[1].text = c, comp

    # Section 8
    doc.add_heading('8. Project Structure', level=1)
    structure = """age-calculator/
├── public/
├── src/
│   ├── components/
│   │   ├── Header.jsx
│   │   ├── DateInput.jsx
│   │   ├── AgeResult.jsx
│   │   ├── AgeBreakdown.jsx
│   │   ├── BirthdayCountdown.jsx
│   │   ├── FunFacts.jsx
│   │   ├── HistoryItem.jsx
│   │   ├── HistoryList.jsx
│   │   └── Footer.jsx
│   ├── utils/
│   │   └── ageCalculator.js
│   ├── App.jsx
│   ├── App.css
│   └── main.jsx
├── index.html
├── package.json
└── vite.config.js"""
    doc.add_paragraph(structure)


    # Section 9
    doc.add_heading('9. Screenshots', level=1)
    
    doc.add_heading('9.1 Initial State (Dark Theme)', level=2)
    doc.add_paragraph('The application loads with a dark theme by default, showing the date input card and empty calculation history.')
    if os.path.exists('screenshots/initial_state.png'):
        doc.add_picture('screenshots/initial_state.png', width=Inches(2.5))

    doc.add_heading('9.2 After Calculation', level=2)
    doc.add_paragraph('After entering a date of birth (May 15, 2003), the application displays:\n- Age: 23 Years, 3 Months, 12 Days\n- Detailed Breakdown: 279 months, 1,215 weeks, 8,505 days, 204,127 hours, 12,247,639 minutes\n- Next Birthday: 261 days to go (Saturday, 15 May 2027)\n- Fun Facts: Born on Thursday, Taurus ♉, Spring, Not a Leap Year baby, Gen Z\n- History entry added')
    if os.path.exists('screenshots/calculation_results.png'):
        doc.add_picture('screenshots/calculation_results.png', width=Inches(2.5))

    doc.add_heading('9.3 Light Theme', level=2)
    doc.add_paragraph('Theme toggle switches seamlessly to a clean light mode with all elements remaining readable and properly themed.')
    if os.path.exists('screenshots/light_theme.png'):
        doc.add_picture('screenshots/light_theme.png', width=Inches(2.5))

    # Section 10
    doc.add_heading('10. Conclusion', level=1)
    doc.add_paragraph('The Phase-I (FA-1) of the Age Calculator mini project has been successfully completed. All core React concepts — JSX, functional components, props, useState, event handling, conditional rendering, and lists & keys — have been implemented across 10 modular components. The application features a responsive, premium UI design with dark/light theme support, input validation, and comprehensive age calculation functionality.')

    doc.add_paragraph('\n\nSubmitted by:\nMagar Dilip Balu (PRN: 125M1H055)\nSYMCA, Semester III\nPimpri Chinchwad College of Engineering')

    doc.save('Progress_Report_FA1.docx')
    print("Successfully generated Progress_Report_FA1.docx")

if __name__ == '__main__':
    main()
